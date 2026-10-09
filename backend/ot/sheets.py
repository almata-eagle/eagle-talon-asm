"""Spreadsheet in, spreadsheet out.

Import: read any checklist (.xlsx or .csv), find the header row, guess which
column is which (EN and JA headings), and let the person confirm the mapping.
Every original cell is kept on the item, so the export can reproduce the
client's own layout with the current answers filled in, plus Talon's extra
columns, a POA&M sheet and a summary.

Spreadsheet content is untrusted: cells are read as values (never formulas),
capped in size, and written back as plain text so nothing becomes a formula
when the client opens the export.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from typing import Any, Optional

from . import db

MAX_ROWS = 5000
MAX_COLS = 60
MAX_CELL = 4000
MAX_UPLOAD = 15 * 1024 * 1024

FIELDS = ("ref", "control", "requirement", "category", "responsible", "level", "status", "owner",
          "notes", "finding", "remediation", "severity", "due_date")

# Header keywords, most specific first. Matched case-insensitively on the heading text.
_GUESS: list[tuple[str, str]] = [
    ("ref", r"\bcci\b|cci[-\s]?id|^id$|^no\.?$|^#$|item\s*(no|#|id)|^ref|番号|項番|ＩＤ"),
    ("control", r"control|nist|800-53|^ap\s?acronym|管理策|コントロール|統制"),
    ("requirement", r"requirement|description|definition|text|title|question|要件|内容|説明|確認項目|チェック項目"),
    ("category", r"categor|^type$|cci type|区分|分類|カテゴリ"),
    ("responsible", r"responsib|designer|^party|implement(ed)?\s*by|責任|担当区分"),
    ("level", r"^level|ufc level|レベル|階層"),
    ("status", r"status|complian|result|implemented\?|^y/n|met\?|状態|結果|適合|実施状況|判定"),
    ("owner", r"assignee|assessor|reviewer|^owner|担当者|確認者"),
    ("finding", r"finding|issue|deficien|weakness|指摘|不備|所見"),
    ("remediation", r"remediat|corrective|mitigat|recommend|action|是正|対策|改善"),
    ("severity", r"severity|risk|priority|criticality|重要度|リスク|優先"),
    ("due_date", r"due|deadline|target date|期限|予定日"),
    ("notes", r"note|comment|remark|observation|evidence|備考|コメント|メモ|証跡"),
]

_STATUS_MAP = [
    ("not_implemented", r"^(no|n|fail(ed)?|non[-\s]?complian\w*|not\s+(implemented|met|compliant)|open|未実施|不適合|未対応|×|✗|ng)$"),
    ("partial", r"^(partial(ly)?(\s+\w+)?|some|in\s+progress|一部\w*|部分\w*|△)$"),
    ("implemented", r"^(yes|y|pass(ed)?|complian\w*|implemented|met|closed|done|ok|実施済\w*|適合|済|○|◯|✓|✔)$"),
    ("inherited", r"^(inherit\w*|継承)"),
    ("na", r"^(n/?a|not\s+applicable|対象外|該当なし|－|-)$"),
    ("impractical", r"^(impractical|infeasible|not\s+feasible|実施困難)"),
    ("not_assessed", r"^(not\s+assessed|tbd|pending|未確認|未評価)?$"),
]
_CATEGORY_MAP = [
    ("dod_defined", r"dod[-\s]?defined"), ("non_designer", r"non[-\s]?designer"), ("designer", r"designer"),
    ("platform_enclave", r"platform|enclave"), ("impractical", r"impractical"),
]


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10] if isinstance(v, dt.date) and not isinstance(v, dt.datetime) else v.isoformat(sep=" ")[:19]
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return (db.clean(v, MAX_CELL) or "")


def read_table(data: bytes, filename: str) -> list[dict]:
    """Every sheet as rows of strings (capped)."""
    if len(data) > MAX_UPLOAD:
        raise ValueError("file is larger than 15 MB")
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = [[_cell(c) for c in r[:MAX_COLS]] for _, r in zip(range(MAX_ROWS), csv.reader(io.StringIO(text)))]
        return [{"name": "CSV", "rows": rows}]
    if not name.endswith((".xlsx", ".xlsm")):
        raise ValueError("upload an .xlsx or .csv file")
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — any parse failure is a bad file
        raise ValueError(f"couldn't read the workbook: {type(e).__name__}") from None
    sheets = []
    for ws in wb.worksheets[:20]:
        rows = []
        for r in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True):
            rows.append([_cell(c) for c in r])
        while rows and not any(rows[-1]):
            rows.pop()
        # read-only mode pads every row to max_col; keep only columns in use.
        width = max((max((j + 1 for j, c in enumerate(r) if c), default=0) for r in rows), default=0)
        rows = [r[:width] for r in rows]
        sheets.append({"name": db.clean(ws.title, 100) or "Sheet", "rows": rows})
    wb.close()
    return sheets


def guess_header_row(rows: list[list[str]]) -> int:
    best, best_score = 0, -1
    for i, r in enumerate(rows[:25]):
        filled = [c for c in r if c]
        if len(filled) < 2:
            continue
        hits = sum(1 for c in filled for _, rx in _GUESS if re.search(rx, c.strip(), re.I))
        score = hits * 3 + len(filled) - sum(1 for c in filled if len(c) > 80) * 2
        if score > best_score:
            best, best_score = i, score
    return best


_PREFER = {"ref": r"\bcci\b|cci[-\s]?id"}   # a CCI column beats a plain row number


def guess_mapping(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for field, rx in _PREFER.items():
        for i, h in enumerate(headers):
            if h and re.search(rx, h.strip(), re.I):
                mapping[field] = i
                used.add(i)
                break
    for field, rx in _GUESS:
        if field in mapping:
            continue
        for i, h in enumerate(headers):
            if i in used or not h:
                continue
            if re.search(rx, h.strip(), re.I):
                mapping[field] = i
                used.add(i)
                break
    return mapping


def preview(data: bytes, filename: str) -> dict:
    sheets = read_table(data, filename)
    out = []
    for s in sheets:
        h = guess_header_row(s["rows"])
        headers = s["rows"][h] if s["rows"] else []
        out.append({"name": s["name"], "header_row": h, "headers": headers, "rows_total": max(0, len(s["rows"]) - h - 1),
                    "sample": s["rows"][h + 1:h + 8], "top": s["rows"][:min(len(s["rows"]), h + 1)][-6:],
                    "mapping": guess_mapping(headers)})
    return {"filename": db.clean(filename, 200), "sheets": out}


def norm_status(v: Optional[str]) -> str:
    v = (v or "").strip().lower()
    for code, rx in _STATUS_MAP:
        if re.search(rx, v, re.I):
            return code
    return "not_assessed"


def norm_category(v: Optional[str]) -> Optional[str]:
    v = (v or "").strip().lower()
    for code, rx in _CATEGORY_MAP:
        if re.search(rx, v, re.I):
            return code
    return None


def norm_level(v: Optional[str]) -> Optional[str]:
    m = re.search(r"([0-5])", v or "")
    return m.group(1) if m else None


def norm_severity(v: Optional[str]) -> Optional[str]:
    v = (v or "").strip().lower()
    for s in ("critical", "high", "moderate", "low"):
        if s in v or (s == "moderate" and "medium" in v):
            return s
    return {"高": "high", "中": "moderate", "低": "low", "緊急": "critical"}.get(v.strip(), None)


def import_rows(con, system_id: int, filename: str, sheet: dict, header_row: int, mapping: dict[str, int],
                actor: str) -> int:
    headers = sheet["rows"][header_row] if sheet["rows"] else []
    mapping = {f: int(i) for f, i in mapping.items() if f in FIELDS and i is not None and 0 <= int(i) < len(headers)}
    rows = [r for r in sheet["rows"][header_row + 1:] if any(r)]

    def get(r, f):
        i = mapping.get(f)
        return r[i] if i is not None and i < len(r) else None

    base = con.execute("SELECT COALESCE(max(sort), 0) FROM ot_items WHERE system_id = ?", (system_id,)).fetchone()[0]
    n = 0
    for k, r in enumerate(rows):
        r = (r + [""] * len(headers))[:max(len(headers), len(r))]
        con.execute(
            "INSERT INTO ot_items (system_id, sort, ref, control, requirement, category, responsible, level, status, "
            "owner, notes, finding, remediation, severity, due_date, source_row, updated_at, updated_by) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (system_id, base + k + 1, db.clean(get(r, "ref"), 100), db.clean(get(r, "control"), 200),
             db.clean(get(r, "requirement"), MAX_CELL), norm_category(get(r, "category")) or db.clean(get(r, "category"), 60),
             db.clean(get(r, "responsible"), 100), norm_level(get(r, "level")), norm_status(get(r, "status")),
             db.clean(get(r, "owner"), 100), db.clean(get(r, "notes"), MAX_CELL), db.clean(get(r, "finding"), MAX_CELL),
             db.clean(get(r, "remediation"), MAX_CELL), norm_severity(get(r, "severity")), db.clean(get(r, "due_date"), 20),
             json.dumps(r, ensure_ascii=False), db.now_iso(), actor))
        n += 1
    template = {"filename": db.clean(filename, 200), "sheet": sheet["name"], "header_row": header_row,
                "headers": headers, "top": sheet["rows"][:header_row], "mapping": mapping,
                "imported_at": db.now_iso(), "imported_by": actor}
    con.execute("UPDATE ot_systems SET template = ?, updated_at = ? WHERE id = ?",
                (json.dumps(template, ensure_ascii=False), db.now_iso(), system_id))
    db.audit(con, actor, "items.import", "system", system_id, {"rows": n, "file": filename, "sheet": sheet["name"],
                                                               "mapping": mapping})
    return n


# ------------------------------------------------------------------ export

STATUS_LABEL = {"not_assessed": "Not assessed", "implemented": "Implemented", "partial": "Partially implemented",
                "not_implemented": "Not implemented", "inherited": "Inherited", "na": "Not applicable",
                "impractical": "Impractical"}
SEV_LABEL = {"low": "Low", "moderate": "Moderate", "high": "High", "critical": "Critical"}
FIELD_LABEL = {"ref": "Ref / CCI", "control": "Control", "requirement": "Requirement", "category": "CCI category",
               "responsible": "Responsible", "level": "UFC level", "status": "Status", "owner": "Owner",
               "notes": "Notes", "finding": "Finding", "remediation": "Remediation", "severity": "Severity",
               "due_date": "Due date"}


_NORMALISED = {"category": norm_category, "level": norm_level, "severity": norm_severity}


def _safe(ws, row: int, col: int, value: Any):
    """Write text so Excel can't treat it as a formula (=, +, -, @ …)."""
    c = ws.cell(row=row, column=col, value=value if value not in (None, "") else None)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        c.data_type = "s"
    return c


def export_workbook(con, engagement_id: int) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    e = db.row(con.execute("SELECT * FROM ot_engagements WHERE id = ?", (engagement_id,)).fetchone())
    if not e:
        raise KeyError("engagement")
    systems = [db.row(r) for r in con.execute("SELECT * FROM ot_systems WHERE engagement_id = ? ORDER BY id", (engagement_id,))]
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    marking = e.get("marking") or ""

    def mark(ws):
        if marking:
            ws.oddHeader.center.text = marking
            ws.oddFooter.center.text = marking

    r = 1
    for label, val in (("Marking", marking), ("Engagement", e["name"]), ("Client", e.get("client")), ("Site", e.get("site")),
                       ("Contract", e.get("contract_ref")), ("Exported", db.now_iso()), ("Tool", "Eagle Talon OT")):
        _safe(summary, r, 1, label).font = bold
        _safe(summary, r, 2, val)
        r += 1
    r += 1
    cols = ["System", "Type", "Phase", "C", "I", "A", *[STATUS_LABEL[s] for s in db.STATUSES], "Total"]
    for j, h in enumerate(cols, 1):
        c = _safe(summary, r, j, h)
        c.font, c.fill = bold, head_fill
    poam: list[list] = []
    used_names = {"Summary", "POA&M"}
    for s in systems:
        items = [db.row(x) for x in con.execute("SELECT * FROM ot_items WHERE system_id = ? ORDER BY sort, id", (s["id"],))]
        ev = {x["item_id"]: (x["n"], x["h"]) for x in con.execute(
            "SELECT item_id, count(*) AS n, group_concat(sha256, ' ') AS h FROM ot_evidence "
            "WHERE item_id IN (SELECT id FROM ot_items WHERE system_id = ?) GROUP BY item_id", (s["id"],))}
        counts = {k: 0 for k in db.STATUSES}
        for it in items:
            counts[it["status"] if it["status"] in counts else "not_assessed"] += 1
            if it["status"] in db.OPEN_STATUSES:
                poam.append([s["name"], it["ref"], it["control"], it["finding"] or it["requirement"], SEV_LABEL.get(it["severity"]),
                             it["remediation"], it["owner"], it["due_date"], STATUS_LABEL.get(it["status"])])
        r += 1
        for j, v in enumerate([s["name"], s.get("system_type"), s.get("phase"), s.get("impact_c"), s.get("impact_i"),
                               s.get("impact_a"), *[counts[k] for k in db.STATUSES], len(items)], 1):
            _safe(summary, r, j, v)

        # One sheet per system, in the imported layout when there is one.
        title = re.sub(r"[\[\]\*\?/\\:]", " ", s["name"])[:28] or f"System {s['id']}"
        base, n = title, 2
        while title in used_names:
            title, n = f"{base[:25]} {n}", n + 1
        used_names.add(title)
        ws = wb.create_sheet(title)
        mark(ws)
        tpl = json.loads(s["template"]) if s.get("template") else None
        # The sheet's own words for each status (○/×/△, Yes/No…), so a changed answer reads like the rest.
        vocab: dict[str, str] = {}
        si = (tpl or {}).get("mapping", {}).get("status")
        if si is not None:
            for it in items:
                src = json.loads(it["source_row"]) if it.get("source_row") else []
                w = src[int(si)] if int(si) < len(src) else ""
                if w:
                    vocab.setdefault(norm_status(w), w)
        extra = ["Talon status", "Finding", "Remediation", "Owner", "Severity", "Due date", "Evidence files",
                 "Evidence SHA-256", "Last updated", "Updated by"]
        if tpl:
            top = tpl.get("top") or []
            for i, trow in enumerate(top, 1):
                for j, v in enumerate(trow, 1):
                    _safe(ws, i, j, v)
            hr = len(top) + 1
            headers = tpl.get("headers") or []
            mapping = {f: int(i) for f, i in (tpl.get("mapping") or {}).items()}
        else:
            hr, headers, mapping = 1, [FIELD_LABEL[f] for f in FIELDS], {f: i for i, f in enumerate(FIELDS)}
            extra = ["Evidence files", "Evidence SHA-256", "Last updated", "Updated by"]
        width = len(headers)
        for j, h in enumerate([*headers, *extra], 1):
            c = _safe(ws, hr, j, h)
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for k, it in enumerate(items, 1):
            rr = hr + k
            cells = json.loads(it["source_row"]) if (tpl and it.get("source_row")) else [None] * width
            cells = (cells + [None] * width)[:width]
            for f, i in mapping.items():
                if i >= width:
                    continue
                cur, orig = it.get(f), cells[i]
                # Keep the client's own wording while it still means the same thing.
                if f == "status":
                    cur = orig if (orig and norm_status(orig) == it["status"]) else vocab.get(it["status"], STATUS_LABEL.get(it["status"]))
                elif f in _NORMALISED and orig and _NORMALISED[f](orig) == cur:
                    cur = orig
                cells[i] = cur
            for j, v in enumerate(cells, 1):
                _safe(ws, rr, j, v)
            n_ev, hashes = ev.get(it["id"], (0, ""))
            tail = ([STATUS_LABEL.get(it["status"]), it["finding"], it["remediation"], it["owner"], SEV_LABEL.get(it["severity"]),
                     it["due_date"]] if tpl else []) + [n_ev or 0, hashes or None, it["updated_at"], it["updated_by"]]
            for j, v in enumerate(tail, width + 1):
                _safe(ws, rr, j, v)
        ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    pw = wb.create_sheet("POA&M")
    mark(pw)
    mark(summary)
    for j, h in enumerate(["System", "Ref / CCI", "Control", "Weakness / finding", "Severity", "Remediation", "Owner",
                           "Due date", "Status"], 1):
        c = _safe(pw, 1, j, h)
        c.font, c.fill = bold, head_fill
    for i, p in enumerate(poam, 2):
        for j, v in enumerate(p, 1):
            _safe(pw, i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
