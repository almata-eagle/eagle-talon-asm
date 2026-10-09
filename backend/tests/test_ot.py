"""Tests for Talon OT (backend/ot): accounts and roles, checklist import and
export, evidence, audit. Run:  cd backend && python -m pytest -q tests
"""
import csv
import hashlib
import io
import json
import sys
from pathlib import Path

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from ot import api as ot_api  # noqa: E402
from ot import auth, db, sheets  # noqa: E402

PW = "correct horse battery"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("OT_DATA_DIR", str(tmp_path / "ot"))
    monkeypatch.delenv("OT_SETUP_CODE", raising=False)
    auth._attempts.clear()
    ot_api.init()
    app = FastAPI()
    app.include_router(ot_api.router)
    return TestClient(app)


def H(tok):
    return {"Authorization": f"Bearer {tok}"}


@pytest.fixture
def admin(client):
    r = client.post("/api/ot/setup", json={"username": "eddy", "password": PW})
    assert r.status_code == 200, r.text
    return r.json()["token"]


def make_user(client, admin, name, role):
    assert client.post("/api/ot/users", headers=H(admin), json={"username": name, "password": PW, "role": role}).status_code == 200
    return client.post("/api/ot/login", json={"username": name, "password": PW}).json()["token"]


def xlsx_bytes(rows, title="Checklist"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CHECKLIST = [
    ["FRCS Cybersecurity Checklist — Building 1234", None, None, None, None],
    ["Prepared for: Example Installation", None, None, None, None],
    [None, None, None, None, None],
    ["CCI", "Control", "Requirement", "Category", "Status", "Comments"],
    ["CCI-000015", "AC-2", "Use automated mechanisms to support account management.", "Designer", "Yes", ""],
    ["CCI-000366", "CM-6", "Implement the security configuration settings.", "Non-Designer", "No", "=HYPERLINK(\"http://x\")"],
    ["CCI-001453", "AC-17(2)", "Protect remote sessions with cryptography.", "DoD-Defined", "Partial", "+cmd|' /C calc'!A0"],
    ["CCI-002418", "SC-8", "Protect transmitted information.", "Platform Enclave", "N/A", "@SUM(1)"],
]


# ------------------------------------------------------------------ accounts

def test_setup_once_then_login(client):
    st = client.get("/api/ot/status").json()
    assert st["users_exist"] is False and "implemented" in st["statuses"]
    r = client.post("/api/ot/setup", json={"username": "eddy", "password": PW})
    assert r.status_code == 200 and r.json()["role"] == "admin"
    assert client.post("/api/ot/setup", json={"username": "mal", "password": PW}).status_code == 409
    assert client.get("/api/ot/status").json()["users_exist"] is True
    tok = client.post("/api/ot/login", json={"username": "eddy", "password": PW}).json()["token"]
    assert client.get("/api/ot/me", headers=H(tok)).json() == {"username": "eddy", "role": "admin"}


def test_setup_code_required_when_set(client, monkeypatch):
    monkeypatch.setenv("OT_SETUP_CODE", "site-7731")
    assert client.get("/api/ot/status").json()["setup_code_required"] is True
    assert client.post("/api/ot/setup", json={"username": "eddy", "password": PW}).status_code == 403
    assert client.post("/api/ot/setup", json={"username": "eddy", "password": PW, "code": "nope"}).status_code == 403
    assert client.post("/api/ot/setup", json={"username": "eddy", "password": PW, "code": "site-7731"}).status_code == 200


def test_password_rules_and_hashing(client):
    assert client.post("/api/ot/setup", json={"username": "eddy", "password": "short"}).status_code == 422
    assert client.post("/api/ot/setup", json={"username": "a b", "password": PW}).status_code == 422
    h = auth.hash_password(PW)
    assert PW not in h and h.startswith("pbkdf2_sha256$")
    assert auth.check_password(PW, h) and not auth.check_password(PW + "x", h)
    assert not auth.check_password(PW, "garbage")


def test_tokens_stored_hashed_and_logout(client, admin):
    con = db.connect()
    stored = [r[0] for r in con.execute("SELECT token_hash FROM ot_sessions")]
    con.close()
    assert admin not in stored and hashlib.sha256(admin.encode()).hexdigest() in stored
    assert client.post("/api/ot/logout", headers=H(admin)).status_code == 200
    assert client.get("/api/ot/me", headers=H(admin)).status_code == 401


def test_login_required_and_throttled(client, admin):
    assert client.get("/api/ot/engagements").status_code == 401
    assert client.get("/api/ot/engagements", headers=H("made-up")).status_code == 401
    for _ in range(auth.MAX_ATTEMPTS):
        assert client.post("/api/ot/login", json={"username": "eddy", "password": "wrong password!"}).status_code == 401
    assert client.post("/api/ot/login", json={"username": "eddy", "password": PW}).status_code == 429


def test_roles(client, admin):
    viewer = make_user(client, admin, "vic", "viewer")
    assessor = make_user(client, admin, "ann", "assessor")
    assert client.post("/api/ot/engagements", headers=H(viewer), json={"name": "X"}).status_code == 403
    eid = client.post("/api/ot/engagements", headers=H(assessor), json={"name": "Bldg 1234"}).json()["id"]
    assert client.get(f"/api/ot/engagements/{eid}", headers=H(viewer)).status_code == 200
    assert client.delete(f"/api/ot/engagements/{eid}", headers=H(assessor)).status_code == 403
    assert client.get("/api/ot/users", headers=H(assessor)).status_code == 403
    assert client.get("/api/ot/audit", headers=H(assessor)).status_code == 403
    # Disabling a user ends their sessions.
    assert client.patch("/api/ot/users/ann", headers=H(admin), json={"disabled": True}).status_code == 200
    assert client.get("/api/ot/me", headers=H(assessor)).status_code == 401
    # An admin can't lock themselves out.
    assert client.patch("/api/ot/users/eddy", headers=H(admin), json={"role": "viewer"}).status_code == 422
    assert client.patch("/api/ot/users/eddy", headers=H(admin), json={"disabled": True}).status_code == 422


# ------------------------------------------------------------------ spreadsheet reading

def test_guess_header_and_mapping():
    tables = sheets.read_table(xlsx_bytes(CHECKLIST), "list.xlsx")
    rows = tables[0]["rows"]
    h = sheets.guess_header_row(rows)
    assert h == 3
    m = sheets.guess_mapping(rows[h])
    assert m == {"ref": 0, "control": 1, "requirement": 2, "category": 3, "status": 4, "notes": 5}


def test_guess_prefers_cci_for_ref():
    assert sheets.guess_mapping(["項番", "CCI", "要件"])["ref"] == 1


def test_guess_japanese_headers():
    m = sheets.guess_mapping(["項番", "要件", "区分", "実施状況", "指摘事項", "是正措置", "担当者", "期限", "備考"])
    assert m == {"ref": 0, "requirement": 1, "category": 2, "status": 3, "finding": 4, "remediation": 5,
                 "owner": 6, "due_date": 7, "notes": 8}


@pytest.mark.parametrize("raw,code", [
    ("Yes", "implemented"), ("Compliant", "implemented"), ("○", "implemented"), ("実施済み", "implemented"),
    ("No", "not_implemented"), ("Not Met", "not_implemented"), ("×", "not_implemented"), ("不適合", "not_implemented"),
    ("Partial", "partial"), ("Partially Implemented", "partial"), ("△", "partial"),
    ("N/A", "na"), ("対象外", "na"), ("Inherited", "inherited"), ("Impractical", "impractical"),
    ("", "not_assessed"), ("TBD", "not_assessed"), ("who knows", "not_assessed"),
])
def test_norm_status(raw, code):
    assert sheets.norm_status(raw) == code


def test_norm_other_fields():
    assert sheets.norm_category("Non-Designer") == "non_designer"
    assert sheets.norm_category("Designer") == "designer"
    assert sheets.norm_category("DoD Defined") == "dod_defined"
    assert sheets.norm_level("Level 3 (FPOC)") == "3"
    assert sheets.norm_severity("Medium") == "moderate" and sheets.norm_severity("高") == "high"


def test_csv_and_bad_files():
    buf = io.StringIO()
    csv.writer(buf).writerows([["ID", "Requirement", "Status"], ["1", "Change default passwords", "No"]])
    t = sheets.read_table(buf.getvalue().encode("utf-8-sig"), "x.csv")
    assert t[0]["rows"][1] == ["1", "Change default passwords", "No"]
    with pytest.raises(ValueError):
        sheets.read_table(b"not a workbook", "x.xlsx")
    with pytest.raises(ValueError):
        sheets.read_table(b"x", "x.exe")
    with pytest.raises(ValueError):
        sheets.read_table(b"x" * (sheets.MAX_UPLOAD + 1), "x.csv")


def test_cells_are_cleaned():
    t = sheets.read_table("Ref,Req\n1,ok\x07bad\x1b[31m\n".encode(), "x.csv")
    assert "\x07" not in t[0]["rows"][1][1] and "\x1b" not in t[0]["rows"][1][1]


# ------------------------------------------------------------------ import → edit → export

def _engagement_with_import(client, tok, rows=CHECKLIST):
    eid = client.post("/api/ot/engagements", headers=H(tok),
                      json={"name": "Bldg 1234 HVAC", "client": "Example", "marking": "CUI"}).json()["id"]
    sid = client.post(f"/api/ot/engagements/{eid}/systems", headers=H(tok),
                      json={"name": "HVAC / BAS", "system_type": "hvac_bas", "phase": "concept",
                            "impact_c": "low", "impact_i": "moderate", "impact_a": "moderate"}).json()["id"]
    prev = client.post("/api/ot/import/preview", headers=H(tok),
                       files={"file": ("checklist.xlsx", xlsx_bytes(rows), "application/octet-stream")})
    assert prev.status_code == 200, prev.text
    p = prev.json()
    sh = p["sheets"][0]
    assert sh["header_row"] == 3 and sh["rows_total"] == 4
    r = client.post(f"/api/ot/systems/{sid}/import", headers=H(tok),
                    json={"upload_id": p["upload_id"], "sheet": sh["name"], "header_row": sh["header_row"],
                          "mapping": sh["mapping"]})
    assert r.status_code == 200 and r.json()["imported"] == 4, r.text
    return eid, sid, p["upload_id"]


def test_import_items_and_upload_is_single_use(client, admin):
    eid, sid, uid = _engagement_with_import(client, admin)
    items = client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"]
    assert [i["status"] for i in items] == ["implemented", "not_implemented", "partial", "na"]
    assert [i["category"] for i in items] == ["designer", "non_designer", "dod_defined", "platform_enclave"]
    assert items[0]["source"]["CCI"] == "CCI-000015"
    again = client.post(f"/api/ot/systems/{sid}/import", headers=H(admin),
                        json={"upload_id": uid, "sheet": "Checklist", "header_row": 3, "mapping": {}})
    assert again.status_code == 410
    bad = client.post(f"/api/ot/systems/{sid}/import", headers=H(admin),
                      json={"upload_id": "../../etc/passwd", "sheet": "Checklist", "header_row": 3, "mapping": {}})
    assert bad.status_code == 422


def test_export_round_trip_keeps_layout_and_blocks_formulas(client, admin):
    eid, sid, _ = _engagement_with_import(client, admin)
    items = client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"]
    r = client.patch(f"/api/ot/items/{items[1]['id']}", headers=H(admin),
                     json={"status": "partial", "finding": "=cmd()", "severity": "high", "remediation": "Apply STIG baseline",
                           "owner": "ann", "due_date": "2026-12-01"})
    assert r.status_code == 200 and r.json()["status"] == "partial"
    x = client.get(f"/api/ot/engagements/{eid}/export.xlsx", headers=H(admin))
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "HVAC   BAS", "POA&M"]
    ws = wb["HVAC   BAS"]
    assert ws.cell(1, 1).value.startswith("FRCS Cybersecurity Checklist")     # their title rows kept
    assert [ws.cell(4, j).value for j in range(1, 7)] == CHECKLIST[3]          # their headers kept
    assert ws.cell(5, 5).value == "Yes"                                        # unchanged answer keeps their wording
    assert ws.cell(6, 5).value == "Partial"                                     # changed answer uses the sheet's own word
    assert ws.cell(5, 4).value == "Designer"                                   # category keeps their wording
    # Nothing that came from a person or a spreadsheet is a formula in the export.
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                assert c.data_type != "f", (sheet.title, c.coordinate, c.value)
    assert ws.cell(6, 6).value is None                 # a formula in their file is read as its value, never copied
    assert ws.cell(7, 6).value.startswith("+cmd|")     # text that looks like a formula stays text
    assert ws.cell(6, 8).value == "=cmd()"             # Talon "Finding" column, text not formula
    poam = wb["POA&M"]
    rows = [[c.value for c in r] for r in poam.iter_rows(min_row=2)]
    assert len(rows) == 2                                                    # the partial and the edited one
    assert any(r[3] == "=cmd()" and r[4] == "High" for r in rows)
    assert wb["Summary"]["B1"].value == "CUI" and ws.oddHeader.center.text == "CUI"


def test_export_without_template(client, admin):
    eid = client.post("/api/ot/engagements", headers=H(admin), json={"name": "社内テスト"}).json()["id"]
    sid = client.post(f"/api/ot/engagements/{eid}/systems", headers=H(admin), json={"name": "Fire alarm"}).json()["id"]
    assert client.post(f"/api/ot/systems/{sid}/items", headers=H(admin),
                       json={"ref": "1", "requirement": "Default passwords changed", "status": "not_implemented"}).status_code == 200
    x = client.get(f"/api/ot/engagements/{eid}/export.xlsx", headers=H(admin))
    assert x.status_code == 200 and "filename*=UTF-8''" in x.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb["Fire alarm"].cell(1, 2).value == "Control"
    assert wb["Fire alarm"].cell(2, 7).value == "Not implemented"           # no template: Talon's labels


def test_export_reuses_symbol_vocabulary(client, admin):
    rows = [["項番", "要件", "実施状況"], ["1", "a", "○"], ["2", "b", "×"], ["3", "c", "△"]]
    eid = client.post("/api/ot/engagements", headers=H(admin), json={"name": "JA"}).json()["id"]
    sid = client.post(f"/api/ot/engagements/{eid}/systems", headers=H(admin), json={"name": "BAS"}).json()["id"]
    p = client.post("/api/ot/import/preview", headers=H(admin), files={"file": ("ja.xlsx", xlsx_bytes(rows), "x")}).json()
    sh = p["sheets"][0]
    client.post(f"/api/ot/systems/{sid}/import", headers=H(admin),
                json={"upload_id": p["upload_id"], "sheet": sh["name"], "header_row": sh["header_row"], "mapping": sh["mapping"]})
    it = client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"]
    client.patch(f"/api/ot/items/{it[1]['id']}", headers=H(admin), json={"status": "partial"})
    client.patch(f"/api/ot/items/{it[2]['id']}", headers=H(admin), json={"status": "implemented"})
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/ot/engagements/{eid}/export.xlsx", headers=H(admin)).content))
    assert [wb["BAS"].cell(r, 3).value for r in (2, 3, 4)] == ["○", "△", "○"]


def test_report_lists_open_items_by_severity(client, admin):
    eid, sid, _ = _engagement_with_import(client, admin)
    items = client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"]
    client.patch(f"/api/ot/items/{items[2]['id']}", headers=H(admin), json={"severity": "critical"})
    rep = client.get(f"/api/ot/engagements/{eid}/report", headers=H(admin)).json()
    s = rep["systems"][0]
    assert s["counts"]["implemented"] == 1 and s["counts"]["na"] == 1
    assert [i["severity"] for i in s["open_items"]] == ["critical", None]


def test_item_validation(client, admin):
    eid, sid, _ = _engagement_with_import(client, admin)
    iid = client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"][0]["id"]
    for bad in ({"status": "done-ish"}, {"level": "7"}, {"severity": "huge"}, {"due_date": "next week"}):
        assert client.patch(f"/api/ot/items/{iid}", headers=H(admin), json=bad).status_code == 422, bad
    assert client.post(f"/api/ot/engagements/{eid}/systems", headers=H(admin),
                       json={"name": "x", "impact_c": "extreme"}).status_code == 422
    assert client.post("/api/ot/engagements", headers=H(admin), json={"name": "  "}).status_code == 422


# ------------------------------------------------------------------ evidence

def test_evidence_hash_dedupe_and_integrity(client, admin):
    eid, sid, _ = _engagement_with_import(client, admin)
    ids = [i["id"] for i in client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"]]
    data = b"\x89PNG\r\n\x1a\n fake screenshot"
    sha = hashlib.sha256(data).hexdigest()
    a = client.post(f"/api/ot/items/{ids[0]}/evidence", headers=H(admin),
                    files={"file": ("../../panel 1.png", data, "image/png")}, data={"note": "BAS login page"}).json()
    b = client.post(f"/api/ot/items/{ids[1]}/evidence", headers=H(admin), files={"file": ("same.png", data, "image/png")}).json()
    assert a["sha256"] == b["sha256"] == sha and a["filename"] == "panel 1.png"
    stored = db.evidence_dir() / sha[:2] / sha
    assert stored.read_bytes() == data
    f = client.get(f"/api/ot/evidence/{a['id']}/file", headers=H(admin))
    assert f.content == data and f.headers["x-content-type-options"] == "nosniff"
    assert "sandbox" in f.headers["content-security-policy"]
    # Deleting one record keeps the shared file; deleting the last removes it.
    client.delete(f"/api/ot/evidence/{a['id']}", headers=H(admin))
    assert stored.exists()
    client.delete(f"/api/ot/evidence/{b['id']}", headers=H(admin))
    assert not stored.exists()


def test_evidence_tamper_and_limits(client, admin, monkeypatch):
    eid, sid, _ = _engagement_with_import(client, admin)
    iid = client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"][0]["id"]
    assert client.post(f"/api/ot/items/{iid}/evidence", headers=H(admin),
                       files={"file": ("run.exe", b"MZ", "application/octet-stream")}).status_code == 422
    assert client.post(f"/api/ot/items/{iid}/evidence", headers=H(admin),
                       files={"file": ("page.html", b"<script>", "text/html")}).status_code == 422
    assert client.post(f"/api/ot/items/99999/evidence", headers=H(admin),
                       files={"file": ("a.txt", b"x", "text/plain")}).status_code == 404
    monkeypatch.setattr(ot_api, "MAX_EVIDENCE", 10)
    assert client.post(f"/api/ot/items/{iid}/evidence", headers=H(admin),
                       files={"file": ("a.txt", b"x" * 11, "text/plain")}).status_code == 413
    monkeypatch.setattr(ot_api, "MAX_EVIDENCE", 25 * 1024 * 1024)
    e = client.post(f"/api/ot/items/{iid}/evidence", headers=H(admin), files={"file": ("cfg.txt", b"hostname bas1", "text/plain")}).json()
    f = client.get(f"/api/ot/evidence/{e['id']}/file", headers=H(admin))
    assert "attachment" in f.headers["content-disposition"]                 # non-images never render inline
    (db.evidence_dir() / e["sha256"][:2] / e["sha256"]).write_bytes(b"tampered")
    assert client.get(f"/api/ot/evidence/{e['id']}/file", headers=H(admin)).status_code == 500


# ------------------------------------------------------------------ audit

def test_every_change_is_audited(client, admin):
    eid, sid, _ = _engagement_with_import(client, admin)
    iid = client.get(f"/api/ot/systems/{sid}/items", headers=H(admin)).json()["items"][0]["id"]
    client.patch(f"/api/ot/items/{iid}", headers=H(admin), json={"status": "not_implemented"})
    client.get(f"/api/ot/engagements/{eid}/export.xlsx", headers=H(admin))
    log = client.get("/api/ot/audit", headers=H(admin)).json()["audit"]
    actions = [a["action"] for a in log]
    for want in ("user.create", "login", "engagement.create", "system.create", "items.import", "item.update",
                 "engagement.export"):
        assert want in actions, want
    upd = next(a for a in log if a["action"] == "item.update")
    assert json.loads(upd["detail"])["status"] == {"from": "implemented", "to": "not_implemented"}
    assert all(a["actor"] == "eddy" for a in log)


def test_ot_data_is_separate_from_talon(client, admin, tmp_path):
    assert db.db_path() == tmp_path / "ot" / "ot.db"
    con = db.connect()
    tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    con.close()
    assert tables and all(t.startswith("ot_") or t == "sqlite_sequence" for t in tables)
